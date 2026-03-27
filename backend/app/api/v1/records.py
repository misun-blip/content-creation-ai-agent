"""
创作记录 & 版本管理 API
"""
import csv
import io
from datetime import datetime, date
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.api import deps
from app.models.user import User
from app.models.record import Record, Version
from app.schemas.record import RecordCreate, RecordUpdate, RecordOut, VersionCreate, VersionOut
from app.utils.response import success, paginated

router = APIRouter()


# ─────────────────────────────────────────────────────────────
# 辅助函数
# ─────────────────────────────────────────────────────────────

def get_record_or_404(record_id: int, user_id: int, db: Session) -> Record:
    record = db.query(Record).filter(
        Record.id == record_id,
        Record.user_id == user_id
    ).first()
    if not record:
        raise HTTPException(status_code=404, detail="创作记录不存在")
    return record


# ─────────────────────────────────────────────────────────────
# 创作记录 CRUD
# ─────────────────────────────────────────────────────────────

@router.get("/")
async def list_records(
    keyword: Optional[str] = Query(None, description="标题关键词"),
    platform: Optional[str] = Query(None, description="平台筛选"),
    status: Optional[str] = Query(None, description="状态筛选 draft/published/archived"),
    start_date: Optional[date] = Query(None, description="开始日期 YYYY-MM-DD"),
    end_date: Optional[date] = Query(None, description="结束日期 YYYY-MM-DD"),
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """获取创作记录列表（支持关键词/平台/状态/时间范围检索）"""
    query = db.query(Record).filter(Record.user_id == current_user.id)

    if keyword:
        query = query.filter(Record.title.ilike(f"%{keyword}%"))
    if platform:
        query = query.filter(Record.platform == platform)
    if status and status in ("draft", "published", "archived"):
        query = query.filter(Record.status == status)
    if start_date:
        query = query.filter(Record.created_at >= datetime.combine(start_date, datetime.min.time()))
    if end_date:
        query = query.filter(Record.created_at <= datetime.combine(end_date, datetime.max.time()))

    total = query.count()
    records = (
        query.order_by(Record.created_at.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )

    items = []
    for r in records:
        items.append({
            "id": r.id,
            "title": r.title,
            "platform": r.platform,
            "status": getattr(r, "status", "draft"),
            "created_at": r.created_at.isoformat() if r.created_at else None,
            "updated_at": r.updated_at.isoformat() if r.updated_at else None,
            "version_count": len(r.versions),
        })

    return paginated(items=items, total=total)


@router.post("/")
async def create_record(
    record_in: RecordCreate,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """创建创作记录，并自动创建第一个版本"""
    status = (record_in.status or "draft").strip() or "draft"
    if status not in ("draft", "published", "archived"):
        status = "draft"
    record = Record(
        user_id=current_user.id,
        title=record_in.title,
        platform=record_in.platform,
        content=record_in.content,
        status=status,
    )
    db.add(record)
    db.flush()  # 获取 id

    if record_in.content:
        v1 = Version(
            record_id=record.id,
            version_number=1,
            content=record_in.content,
            change_note="初始版本",
            is_ai_generated=True,
            created_by=current_user.id,
        )
        db.add(v1)

    db.commit()
    db.refresh(record)
    return success(data=RecordOut.model_validate(record).model_dump(), message="创建成功")


@router.get("/{record_id}")
async def get_record(
    record_id: int,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """获取单条创作记录详情（含所有版本）"""
    record = get_record_or_404(record_id, current_user.id, db)
    return success(data=RecordOut.model_validate(record).model_dump())


@router.put("/{record_id}")
async def update_record(
    record_id: int,
    record_in: RecordUpdate,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """更新创作记录基本信息（不创建新版本）"""
    record = get_record_or_404(record_id, current_user.id, db)

    if record_in.title is not None:
        record.title = record_in.title
    if record_in.platform is not None:
        record.platform = record_in.platform
    if record_in.content is not None:
        record.content = record_in.content
    if record_in.status is not None:
        if record_in.status in ("draft", "published", "archived"):
            record.status = record_in.status

    db.commit()
    db.refresh(record)
    return success(data=RecordOut.model_validate(record).model_dump(), message="更新成功")


@router.delete("/{record_id}")
async def delete_record(
    record_id: int,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """删除创作记录（级联删除所有版本）"""
    record = get_record_or_404(record_id, current_user.id, db)
    db.delete(record)
    db.commit()
    return success(message="删除成功")


# ─────────────────────────────────────────────────────────────
# 版本管理
# ─────────────────────────────────────────────────────────────

@router.get("/{record_id}/versions")
async def list_versions(
    record_id: int,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """获取某条记录的所有版本列表"""
    get_record_or_404(record_id, current_user.id, db)
    versions = (
        db.query(Version)
        .filter(Version.record_id == record_id)
        .order_by(Version.version_number)
        .all()
    )
    data = [VersionOut.model_validate(v).model_dump() for v in versions]
    return success(data=data)


@router.post("/{record_id}/versions")
async def create_version(
    record_id: int,
    version_in: VersionCreate,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """为某条记录创建新版本"""
    record = get_record_or_404(record_id, current_user.id, db)

    latest = (
        db.query(Version)
        .filter(Version.record_id == record_id)
        .order_by(Version.version_number.desc())
        .first()
    )
    next_number = (latest.version_number + 1) if latest else 1

    is_ai = getattr(version_in, "is_ai_generated", True)
    version = Version(
        record_id=record_id,
        version_number=next_number,
        content=version_in.content,
        change_note=version_in.change_note,
        is_ai_generated=is_ai if isinstance(is_ai, bool) else True,
        created_by=current_user.id,
    )
    db.add(version)

    # 同步更新记录内容为最新版本
    record.content = version_in.content
    db.commit()
    db.refresh(version)
    return success(data=VersionOut.model_validate(version).model_dump(), message="版本创建成功")


@router.get("/{record_id}/versions/{version_id}")
async def get_version(
    record_id: int,
    version_id: int,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """获取指定版本详情"""
    get_record_or_404(record_id, current_user.id, db)
    version = db.query(Version).filter(
        Version.id == version_id,
        Version.record_id == record_id
    ).first()
    if not version:
        raise HTTPException(status_code=404, detail="版本不存在")
    return success(data=VersionOut.model_validate(version).model_dump())


@router.post("/{record_id}/versions/{version_id}/restore")
async def restore_version(
    record_id: int,
    version_id: int,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """回滚到指定版本（创建一个新版本，内容来自目标版本）"""
    record = get_record_or_404(record_id, current_user.id, db)
    target = db.query(Version).filter(
        Version.id == version_id,
        Version.record_id == record_id
    ).first()
    if not target:
        raise HTTPException(status_code=404, detail="版本不存在")

    latest = (
        db.query(Version)
        .filter(Version.record_id == record_id)
        .order_by(Version.version_number.desc())
        .first()
    )
    next_number = (latest.version_number + 1) if latest else 1

    restored = Version(
        record_id=record_id,
        version_number=next_number,
        content=target.content,
        change_note=f"回滚自版本 v{target.version_number}",
        is_ai_generated=getattr(target, "is_ai_generated", True),
        created_by=current_user.id,
    )
    db.add(restored)
    record.content = target.content
    db.commit()
    db.refresh(restored)
    return success(data=VersionOut.model_validate(restored).model_dump(), message="回滚成功")


@router.delete("/{record_id}/versions/{version_id}")
async def delete_version(
    record_id: int,
    version_id: int,
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """删除指定版本（至少保留一个版本）"""
    get_record_or_404(record_id, current_user.id, db)
    version = db.query(Version).filter(
        Version.id == version_id,
        Version.record_id == record_id
    ).first()
    if not version:
        raise HTTPException(status_code=404, detail="版本不存在")

    count = db.query(Version).filter(Version.record_id == record_id).count()
    if count <= 1:
        raise HTTPException(status_code=400, detail="至少需要保留一个版本")

    db.delete(version)
    db.commit()
    return success(message="版本删除成功")


# ─────────────────────────────────────────────────────────────
# 导出
# ─────────────────────────────────────────────────────────────

@router.get("/export/csv")
async def export_records_csv(
    keyword: Optional[str] = Query(None),
    platform: Optional[str] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """导出创作记录为 CSV 文件"""
    query = db.query(Record).filter(Record.user_id == current_user.id)

    if keyword:
        query = query.filter(Record.title.ilike(f"%{keyword}%"))
    if platform:
        query = query.filter(Record.platform == platform)
    if start_date:
        query = query.filter(Record.created_at >= datetime.combine(start_date, datetime.min.time()))
    if end_date:
        query = query.filter(Record.created_at <= datetime.combine(end_date, datetime.max.time()))

    records = query.order_by(Record.created_at.desc()).all()

    from zoneinfo import ZoneInfo
    tz_cn = ZoneInfo("Asia/Shanghai")

    def fmt_time(dt):
        if not dt:
            return ""
        # 若数据库存的是 naive UTC 时间，手动加上 UTC 信息再转北京时间
        if dt.tzinfo is None:
            from datetime import timezone
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(tz_cn).strftime("%Y-%m-%d %H:%M:%S")

    # 用 BytesIO 直接写 UTF-8-BOM 字节流，Excel 打开不乱码
    output = io.BytesIO()
    output.write(b"\xef\xbb\xbf")  # UTF-8 BOM
    text_output = io.TextIOWrapper(output, encoding="utf-8", newline="")
    writer = csv.writer(text_output)
    writer.writerow(["ID", "标题", "平台", "版本数", "创建时间", "更新时间", "内容"])
    for r in records:
        writer.writerow([
            r.id,
            r.title,
            r.platform,
            len(r.versions),
            fmt_time(r.created_at),
            fmt_time(r.updated_at),
            (r.content or "").replace("\n", " "),
        ])
    text_output.flush()
    text_output.detach()  # 解除绑定，避免关闭 BytesIO
    output.seek(0)

    filename = f"records_{datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d%H%M%S')}.csv"
    return StreamingResponse(
        iter([output.read()]),
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _records_export_query(
    keyword: Optional[str],
    platform: Optional[str],
    start_date: Optional[date],
    end_date: Optional[date],
    current_user: User,
    db: Session,
):
    """共用：按条件查询创作记录列表"""
    query = db.query(Record).filter(Record.user_id == current_user.id)
    if keyword:
        query = query.filter(Record.title.ilike(f"%{keyword}%"))
    if platform:
        query = query.filter(Record.platform == platform)
    if start_date:
        query = query.filter(Record.created_at >= datetime.combine(start_date, datetime.min.time()))
    if end_date:
        query = query.filter(Record.created_at <= datetime.combine(end_date, datetime.max.time()))
    return query.order_by(Record.created_at.desc()).all()


@router.get("/export/txt")
async def export_records_txt(
    keyword: Optional[str] = Query(None),
    platform: Optional[str] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """导出创作记录为纯文本"""
    from zoneinfo import ZoneInfo
    tz_cn = ZoneInfo("Asia/Shanghai")

    def fmt_time(dt):
        if not dt:
            return ""
        if dt.tzinfo is None:
            from datetime import timezone
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(tz_cn).strftime("%Y-%m-%d %H:%M:%S")

    records = _records_export_query(keyword, platform, start_date, end_date, current_user, db)
    lines = []
    for r in records:
        lines.append(f"【{r.id}】{r.title}")
        lines.append(f"平台：{r.platform} | 创建：{fmt_time(r.created_at)} | 更新：{fmt_time(r.updated_at)}")
        lines.append((r.content or "").replace("\r\n", "\n"))
        lines.append("\n" + "=" * 60 + "\n")

    content = "\n".join(lines).encode("utf-8")
    filename = f"records_{datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d%H%M%S')}.txt"
    return StreamingResponse(
        iter([content]),
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/docx")
async def export_records_docx(
    keyword: Optional[str] = Query(None),
    platform: Optional[str] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """导出创作记录为 Word 文档"""
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="导出 Word 需要安装 python-docx：pip install python-docx",
        )

    from zoneinfo import ZoneInfo
    tz_cn = ZoneInfo("Asia/Shanghai")

    def fmt_time(dt):
        if not dt:
            return ""
        if dt.tzinfo is None:
            from datetime import timezone
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(tz_cn).strftime("%Y-%m-%d %H:%M:%S")

    records = _records_export_query(keyword, platform, start_date, end_date, current_user, db)
    doc = Document()
    doc.add_heading("创作记录导出", 0)
    for r in records:
        doc.add_heading(f"【{r.id}】{r.title}", level=1)
        p = doc.add_paragraph()
        p.add_run(f"平台：{r.platform} | 创建：{fmt_time(r.created_at)} | 更新：{fmt_time(r.updated_at)}")
        doc.add_paragraph((r.content or "").replace("\r\n", "\n"))
        doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    filename = f"records_{datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d%H%M%S')}.docx"
    return StreamingResponse(
        iter([buf.read()]),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/pdf")
async def export_records_pdf(
    keyword: Optional[str] = Query(None),
    platform: Optional[str] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
):
    """导出创作记录为 PDF"""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="导出 PDF 需要安装 reportlab：pip install reportlab",
        )

    from zoneinfo import ZoneInfo
    tz_cn = ZoneInfo("Asia/Shanghai")

    def fmt_time(dt):
        if not dt:
            return ""
        if dt.tzinfo is None:
            from datetime import timezone
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(tz_cn).strftime("%Y-%m-%d %H:%M:%S")

    records = _records_export_query(keyword, platform, start_date, end_date, current_user, db)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=72)
    styles = getSampleStyleSheet()
    story = [Paragraph("创作记录导出", styles["Title"])]
    for r in records:
        title = f"【{r.id}】{r.title}"
        story.append(Paragraph(title.replace("<", "&lt;").replace(">", "&gt;"), styles["Heading1"]))
        meta = f"平台：{r.platform} | 创建：{fmt_time(r.created_at)} | 更新：{fmt_time(r.updated_at)}"
        story.append(Paragraph(meta.replace("<", "&lt;").replace(">", "&gt;"), styles["Normal"]))
        content = (r.content or "").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
        story.append(Paragraph(content, styles["Normal"]))
        story.append(Spacer(1, 12))
    doc.build(story)
    buf.seek(0)
    filename = f"records_{datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d%H%M%S')}.pdf"
    return StreamingResponse(
        iter([buf.read()]),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
